from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate
from reportlab.platypus import Paragraph
from reportlab.platypus import Spacer

from docx import Document


def export_pdf(filename, title, content):

    document = SimpleDocTemplate(filename)

    styles = getSampleStyleSheet()

    story = []

    story.append(
        Paragraph(title, styles['Heading1'])
    )

    story.append(
        Spacer(1,12)
    )

    for line in content.split("\n"):

        story.append(
            Paragraph(line, styles["BodyText"])
        )

    document.build(story)



def export_docx(filename, title, content):

    document = Document()

    document.add_heading(
        title,
        level=1
    )

    for line in content.split("\n"):

        document.add_paragraph(line)

    document.save(filename)